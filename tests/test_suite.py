"""星火智教 · 零依赖测试套件（赛事版）。

Run
---
    python tests/test_suite.py          # 全量
    python tests/test_suite.py kb llm   # 只跑指定分组

覆盖：config / kb / llm / rural / workflow(mock+llm病态) / assembly /
storage / app。所有断言用真实运行结果验证，失败即非零退出码。
"""

from __future__ import annotations

import importlib
import json
import os
import re
import sqlite3
import sys
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
os.chdir(ROOT)  # 保证相对路径（data/ 等）与从仓库根启动一致

PASS, FAIL = 0, 0
FAILURES: list[str] = []


def check(name: str, condition: bool, detail: str = "") -> None:
    """Record one assertion result.

    Parameters
    ----------
    name : str
        Case description.
    condition : bool
    detail : str
        Extra context printed on failure.

    Returns
    -------
    None
    """
    global PASS, FAIL
    if condition:
        PASS += 1
        print(f"  ✅ {name}")
    else:
        FAIL += 1
        FAILURES.append(f"{name} {detail}".strip())
        print(f"  ❌ {name} {detail}")


def section(title: str) -> None:
    """Print a section header.

    Parameters
    ----------
    title : str

    Returns
    -------
    None
    """
    print(f"\n=== {title} ===")


# ---------------------------------------------------------------------------
# config
# ---------------------------------------------------------------------------
def test_config() -> None:
    section("config 配置层")
    from xhzh import config
    importlib.reload(config)
    check("默认 mock 模式", config.MODE == "mock")
    check("默认端口 7860", config.PORT == 7860)
    check("默认不开启公网分享", config.SHARE is False)

    os.environ["XHZH_PORT"] = "abc"
    importlib.reload(config)
    check("非法 PORT 回退默认值（不崩溃）", config.PORT == 7860)
    os.environ.pop("XHZH_PORT")

    os.environ["XHZH_LLM_TIMEOUT"] = "abc"
    importlib.reload(config)
    check("非法 LLM_TIMEOUT 回退默认值（不崩溃）", config.LLM_TIMEOUT == 90.0)
    os.environ["XHZH_LLM_TIMEOUT"] = "-5"
    importlib.reload(config)
    check("非正 LLM_TIMEOUT 回退默认值", config.LLM_TIMEOUT == 90.0)
    os.environ.pop("XHZH_LLM_TIMEOUT")
    os.environ["XHZH_LLM_MAX_RETRIES"] = "-1"
    importlib.reload(config)
    check("负数 LLM_MAX_RETRIES 归零保护", config.LLM_MAX_RETRIES == 0)
    os.environ.pop("XHZH_LLM_MAX_RETRIES")
    importlib.reload(config)

    os.environ["XHZH_SHARE"] = "1"
    importlib.reload(config)
    check("XHZH_SHARE=1 生效", config.SHARE is True)
    os.environ.pop("XHZH_SHARE")
    importlib.reload(config)

    from xhzh.adp_client import AdpClient
    try:
        AdpClient("bogus-mode")
        check("非法模式应拒绝", False)
    except ValueError:
        check("非法模式应拒绝", True)

    os.environ["XHZH_MODE"] = "ADP"  # 大写应被归一化
    importlib.reload(config)
    check("模式名大小写归一化", config.MODE == "adp")
    os.environ.pop("XHZH_MODE")
    importlib.reload(config)


# ---------------------------------------------------------------------------
# kb 知识库
# ---------------------------------------------------------------------------
def test_kb() -> None:
    section("kb 知识库解析与检索")
    from xhzh import kb

    cards = kb.load_cards()
    check("知识库解析出足量卡片", len(cards) >= 140, f"实际 {len(cards)}")

    hits = kb.retrieve("数学", "七年级", "一元一次方程的实际应用")
    check("命中预期卡片", hits and hits[0].card_id == "MAT-7A-3-0-K01", str([c.card_id for c in hits]))

    hits = kb.retrieve("数学", "九年级", "完全凭空捏造的课题")
    check("无匹配返回空（EARS-3.2 路径）", hits == [])

    hits = kb.retrieve("数学", "七年级", "一个完全虚构的课题名称")
    check("虚构课题不得误命中（相关度阈值）", hits == [], str([c.card_id for c in hits]))

    hits = kb.retrieve("语文", "八年级", "量子计算机的散热原理")
    check("跨学科虚构课题返回空", hits == [])

    hits = kb.retrieve("数学", "七年级", "方程")
    check("超短课题经名称子串命中", hits and hits[0].card_id == "MAT-7A-3-0-K01", str([c.card_id for c in hits]))

    hits = kb.retrieve("数学", "九年级", "函数")
    check("超短课题词面命中（≤2 bigram）", len(hits) > 0, str([c.card_id for c in hits]))

    # C-1 回归：教案环节须保留时长与活动描述（评估报告发现的渲染丢字段）
    from xhzh.workflow import _render_framework
    text = _render_framework({"stages": [{"name": "情境导入", "minutes": 5, "desc": "情境引出问题"}]})
    check("环节保留 minutes/desc（C-1）", "情境导入（5 分钟）：情境引出问题" in text, text)

    hits = kb.retrieve("数学", "七年级", "的与和及 ！！")
    check("纯停停词/标点课题不炸", isinstance(hits, list))

    hits = kb.retrieve("地理", "七年级", "地球运动")
    check("范围外学科返回空", hits == [])

    # 病态 KB：格式垃圾文本不得崩溃
    with tempfile.NamedTemporaryFile("w", suffix=".md", encoding="utf-8", delete=False) as fh:
        fh.write("# 无卡片文档\n\n随便写点### 不完整行\n\n### BAD-ID 名称\n- 概念定义：x\n")
        bad_path = fh.name
    try:
        os.environ["XHZH_KB_PATH"] = bad_path
        kb._CARDS_CACHE = None
        bad_cards = kb.load_cards()
        check("格式错误的 KB 静默跳过（不崩溃）", bad_cards == [])
        hits = kb.retrieve("数学", "七年级", "一元一次方程")
        check("空库检索返回空", hits == [])
    finally:
        os.environ.pop("XHZH_KB_PATH")
        kb._CARDS_CACHE = None
        os.unlink(bad_path)

    # 缺失 KB 文件：明确报错而非静默空库
    os.environ["XHZH_KB_PATH"] = "/nonexistent/kb.md"
    kb._CARDS_CACHE = None
    try:
        kb.load_cards()
        check("缺失 KB 应抛 FileNotFoundError", False)
    except FileNotFoundError:
        check("缺失 KB 应抛 FileNotFoundError", True)
    finally:
        os.environ.pop("XHZH_KB_PATH")
        kb._CARDS_CACHE = None


# ---------------------------------------------------------------------------
# llm
# ---------------------------------------------------------------------------
def test_llm() -> None:
    section("llm 客户端（解析与配置）")
    from xhzh import llm

    check("未配置时 is_configured=False", llm.is_configured() is False)
    try:
        llm.chat([{"role": "user", "content": "hi"}])
        check("未配置调用应报错", False)
    except llm.LLMError as error:
        check("未配置调用应报错", "LLM_BASE_URL" in str(error))

    fenced = '前言\n```json\n{"a": 1, "b": {"c": "}"}}\n```\n后记：见附录 {X}。'
    parsed = llm.extract_json(fenced)
    check("围栏+嵌套+尾部杂文", parsed["b"]["c"] == "}", str(parsed))

    escaped = '{"board": "板书{核心} \\"引号\\" → 步骤"}'
    check("字符串内花括号与转义引号", llm.extract_json(escaped)["board"].startswith("板书"))

    # 字符串内裸英文引号（真实失败样本）：应自动转义修复
    stray = '{"objectives": ["能解决如"配套问题""行程问题"等应用题"]}'
    repaired = llm.extract_json(stray)
    check("裸引号自动修复（JSON 容错）", len(repaired["objectives"]) == 1, str(repaired)[:100])

    # schema 漂移：dict 值应展平为字符串列表
    from xhzh.workflow import _as_str_list
    check("dict 漂移展平", _as_str_list({"basic": ["a", "b"], "advanced": "c"}) == ["a", "b", "c"])

    for bad, label in [
        ("抱歉，我不明白这个问题", "无 JSON"),
        ("[1, 2, 3]", "顶层是数组"),
        ('{"a": 1', "截断（括号不配对）"),
        ("12345", "裸数字"),
    ]:
        try:
            llm.extract_json(bad)
            check(f"{label} 应拒绝", False)
        except llm.LLMError:
            check(f"{label} 应拒绝", True)


# ---------------------------------------------------------------------------
# rural 乡土转换
# ---------------------------------------------------------------------------
def test_rural() -> None:
    section("rural 乡土案例转换")
    from xhzh import rural

    text = rural.ruralize("小明从地铁站出发去购物中心买了 3 支钢笔。")
    check("长词优先（地铁站→乡镇客运站）", "乡镇客运站" in text, text)
    check("购物中心→农贸市场", "农贸市场" in text, text)

    text = rural.ruralize("他坐地铁去镇上游乐园。")
    check("普通地铁词替换", "乡镇班车" in text and "地铁" not in text, text)

    check("无金额文本无告警", rural.check_money_sanity("买 3 支钢笔，每支 8 元") == [])
    warnings = rural.check_money_sanity("在农贸市场批发蔬菜花费 1500 元")
    check("大额金额触发复核", len(warnings) == 1 and "1500" in warnings[0], str(warnings))

    warnings = rural.check_money_sanity("买化肥花了 ３６００ 元（全角数字）")
    check("全角数字金额识别", len(warnings) == 1, str(warnings))

    warnings = rural.check_money_sanity("买化肥花了 3600 块")
    check("口语金额单位“块”识别", len(warnings) == 1, str(warnings))


# ---------------------------------------------------------------------------
# workflow（mock 全链路 + llm 病态输出）
# ---------------------------------------------------------------------------
def _bad_llm_outputs():
    """病态 LLM 输出样本。"""
    return [
        "抱歉，我不太明白。",  # 无 JSON
        '{"standard_ref": "x", "objectives": "理解概念", "stages": "情境导入"}',  # 字段类型漂移
        '{"standard_ref": "x", "objectives": ["a"], "key_points": ["b"], '
        '"difficult_points": ["c"], "stages": [{"name": "导入", "minutes": 5}]}',  # 合法但极简
    ]


def test_workflow_mock() -> None:
    section("workflow mock 全链路")
    from xhzh.workflow import run_generation, render_plan_text

    events = list(run_generation("七年级", "数学", "人教版", "一元一次方程的实际应用"))
    check("末事件为 done", events[-1]["type"] == "done")
    result = events[-1]["result"]
    package = result["package"]
    check("三件套落盘", all(Path(p).exists() for p in package.values()))
    plan = render_plan_text(result)
    check("教案含课标依据", "课标依据" in plan)
    check("教案含双目标", "基础巩固" in plan and "拓展提升" in plan)
    check("教案含乡土化告警区", "随堂练习" in plan)


def test_workflow_llm_pathological() -> None:
    section("workflow llm 病态输出（重试与规整）")
    from xhzh import llm, workflow
    from xhzh.adp_client import AdpClient

    # 场景1：持续返回无 JSON —— 应重试后明确失败，不静默
    calls = {"n": 0}

    def always_garbage(messages, **kwargs):
        calls["n"] += 1
        return "抱歉，我无法回答。"

    with mock.patch.object(llm, "chat", always_garbage), mock.patch.object(llm, "is_configured", lambda: True):
        client = AdpClient("llm")
        try:
            client.call_node(0, {"grade": "七年级", "subject": "数学", "textbook": "人教版", "topic": "一元一次方程"}, {})
            check("持续病态输出应抛错", False)
        except llm.LLMError:
            expected_max = 3  # 1 次首发 + 2 次重试（工作流设计 D2）
            check(f"重试次数符合 D2（{calls['n']} 次，≤{expected_max}）", calls["n"] == expected_max, f"实际 {calls['n']}")

    # 场景2：字段类型漂移（数组→字符串）—— 规整后不崩溃
    def drifted(messages, **kwargs):
        return '{"standard_ref": "x", "objectives": "理解概念", "key_points": "重点", "difficult_points": "难点", "stages": [{"name": "导入", "minutes": 5}]}'

    with mock.patch.object(llm, "chat", drifted), mock.patch.object(llm, "is_configured", lambda: True):
        client = AdpClient("llm")
        out = client.call_node(0, {"grade": "七年级", "subject": "数学", "textbook": "人教版", "topic": "一元一次方程"}, {})
        framework = workflow._render_framework(out)
        check("字符串漂移被规整", "教学目标：1. 理解概念" in framework, framework)

    # 场景3：合法极简输出 —— 全链路可跑通
    def minimal(messages, **kwargs):
        return '{"standard_ref": "x", "objectives": ["a"], "key_points": ["b"], "difficult_points": ["c"], "stages": [{"name": "导入", "minutes": 5}]}'

    with mock.patch.object(llm, "chat", minimal), mock.patch.object(llm, "is_configured", lambda: True):
        events = list(workflow.run_generation("七年级", "数学", "人教版", "一元一次方程", client=AdpClient("llm")))
        check("llm 极简输出端到端可跑", events[-1]["type"] == "done")
        result = events[-1]["result"]
        check("产物落盘", all(Path(p).exists() for p in result["package"].values()))


# ---------------------------------------------------------------------------
# assembly
# ---------------------------------------------------------------------------
def test_assembly() -> None:
    section("assembly 组装层")
    from xhzh.assembly import build_package

    base = {
        "meta": {"grade": "七年级", "subject": "数学", "textbook": "人教版", "topic": "边界探针"},
        "goals_basic": ["a"], "goals_extend": ["b"],
        "exercises_basic": ["c"], "exercises_extend": ["d"], "board": "e",
    }

    paths = [build_package(dict(base))["zip"] for _ in range(3)]
    check("同秒连续组装互不覆盖", len({Path(p).name for p in paths}) == 3, str(paths))

    long_meta = dict(base["meta"], topic="一元一次方程的实际应用" * 30)
    pkg = build_package({**base, "meta": long_meta})
    check("300 字课题不崩溃", Path(pkg["docx"]).exists())

    evil_meta = dict(base["meta"], topic="../../../../etc/passwd")
    pkg = build_package({**base, "meta": evil_meta})
    name = Path(pkg["zip"]).name
    check("路径穿越被过滤", "/" not in name and ".." not in name, name)
    check("产物落在 OUTPUT_DIR", Path(pkg["zip"]).parent == Path("data/packages").resolve())

    # C-2 回归：PPT 首页副标题须含年级/教材（此前 if index==0: continue 导致为空）
    from pptx import Presentation
    prs = Presentation(pkg["pptx"])
    subtitle = prs.slides[0].placeholders[1].text
    check("PPT 首页副标题非空（C-2）", "人教版" in subtitle and "年级" in subtitle, subtitle[:60])

    emoji_meta = dict(base["meta"], topic="几何图形 🚀 认识")
    pkg = build_package({**base, "meta": emoji_meta})
    check("emoji 课题可落盘", Path(pkg["docx"]).exists())

    # 课题含换行（textarea 多行输入）不应破坏 docx 标题
    nl_meta = dict(base["meta"], topic="课题\n副标题")
    pkg = build_package({**base, "meta": nl_meta})
    from docx import Document
    doc = Document(pkg["docx"])
    check("换行课题 docx 仍可读", len(doc.paragraphs) > 0)

    # 学生向课件脚本：模板脚本 + 卡片式渲染 + 版面不越界
    from xhzh.courseware import build_script
    script = build_script(base["meta"], base, "mock")
    check("模板脚本首封面末小结", script[0]["kind"] == "cover" and script[-1]["kind"] == "summary")
    check("脚本页型均为枚举内", all(s["kind"] in ("cover", "hook", "quest", "explore", "route",
                                              "practice", "game", "answer", "board", "summary")
                                 for s in script))
    engaged = build_package({**base, "courseware": script})
    prs2 = Presentation(engaged["pptx"])
    check("学生向 PPT 页数与脚本一致", len(prs2.slides) == len(script), f"{len(prs2.slides)} vs {len(script)}")
    EMU_IN = 914400
    width_in, height_in = prs2.slide_width / EMU_IN, prs2.slide_height / EMU_IN
    check("画布为 16:9", abs(width_in / height_in - 16 / 9) < 0.01, f"{width_in:.2f}x{height_in:.2f}")
    outside = [
        (i + 1, sh.shape_type)
        for i, slide in enumerate(prs2.slides)
        for sh in slide.shapes
        if sh.left < 0 or sh.top < 0 or sh.left + sh.width > prs2.slide_width
        or sh.top + sh.height > prs2.slide_height
    ]
    check("课件形状不越出画布", not outside, str(outside[:3]))
    cover_text = " ".join(
        sh.text_frame.text for sh in prs2.slides[0].shapes if sh.has_text_frame
    )
    check("封面含课题与年级", "边界探针" in cover_text and "七年级" in cover_text, cover_text[:60])
    fallback = build_script(base["meta"], base, "llm")
    check("llm 不可用时降级为模板脚本", len(fallback) > 0)


# ---------------------------------------------------------------------------
# visuals（插画素材库 + 可选生图 + 矢量降级）
# ---------------------------------------------------------------------------
#: 实测道具：6 张素材右下角平台水印「AI生成 WORKBUDDY」的最左/最上起点
#: （像素级检测：x∈[92.3%, 98.7%]W、y∈[94.1%, 98.7%]H）。
#: 该值只用于"守住承诺"，不作为实现参数——实现侧读 visuals.WATERMARK_GUARD。
WATERMARK_MEASURED_START = 0.923

#: 守卫下限：右裁必须 ≥ 1 - 水印起点，否则字头会露进画面（6% 即会泄漏）。
WATERMARK_MIN_GUARD = round(1.0 - WATERMARK_MEASURED_START, 4)


def test_visuals() -> None:
    section("visuals 插画素材库与生图兜底")
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    from xhzh import config, imagegen, visuals
    from xhzh.assembly import build_package
    from xhzh.courseware import build_script

    base = {
        "meta": {"grade": "七年级", "subject": "数学", "textbook": "人教版", "topic": "插画探针"},
        "goals_basic": ["a"], "goals_extend": ["b"],
        "exercises_basic": ["c"], "exercises_extend": ["d"], "board": "e",
    }
    script = build_script(base["meta"], base, "mock")

    # 1) 素材库 6 个主题齐全（jpg 或 png 至少一个）
    missing = [
        theme for theme in visuals.SCENE_THEMES
        if not any((visuals.ASSET_DIR / f"{theme}{suffix}").is_file() for suffix in visuals.IMAGE_SUFFIXES)
    ]
    check("6 个场景主题素材齐备（jpg 或 png）", not missing, str(missing))
    check("场景主题均有中文名", all(visuals.SCENE_LABELS.get(t) for t in visuals.SCENE_THEMES))

    # 2) 素材体积 < 300KB（县乡低带宽）
    oversized = []
    for theme in visuals.SCENE_THEMES:
        image = visuals.scene_image(theme)
        if image is None:
            oversized.append(f"{theme}:缺失")
        elif image.stat().st_size >= 300 * 1024:
            oversized.append(f"{theme}:{image.stat().st_size // 1024}KB")
    check("素材体积均 < 300KB（低带宽）", not oversized, str(oversized))

    # 3) 解析优先级链的边界
    market = visuals.scene_image("market")
    check("scene_image('market') 返回存在的 Path", isinstance(market, Path) and market.is_file(), str(market))
    check("素材库优先命中 .jpg 小图", market is not None and market.suffix == ".jpg", str(market))
    check("未知 theme 返回 None", visuals.scene_image("unknown") is None)
    check("theme 大小写/空格归一化", visuals.scene_image("  MARKET  ") is not None)

    # 4) 生图适配器默认关闭 + 失败静默
    env_keys = (
        "XHZH_IMAGE_BASE_URL", "XHZH_IMAGE_API_KEY", "XHZH_IMAGE_MODEL",
        "XHZH_IMAGE_ALLOW", "XHZH_IMAGE_TIMEOUT",
    )
    saved = {key: os.environ.pop(key, None) for key in env_keys}
    try:
        check("生图未配置时 is_configured=False", imagegen.is_configured() is False)
        check("生图默认关闭（AIGC 须人工审核）", imagegen.is_enabled() is False)
        os.environ["XHZH_IMAGE_BASE_URL"] = "http://127.0.0.1:9/v1"  # 必然拒连，验证失败路径
        os.environ["XHZH_IMAGE_API_KEY"] = "test-key-not-real"
        os.environ["XHZH_IMAGE_MODEL"] = "test-model"
        check("三件套齐备时 is_configured=True", imagegen.is_configured() is True)
        check("有凭据但未显式开启仍不启用", imagegen.is_enabled() is False)
        os.environ["XHZH_IMAGE_ALLOW"] = "1"
        os.environ["XHZH_IMAGE_TIMEOUT"] = "2"
        check("XHZH_IMAGE_ALLOW=1 才启用", imagegen.is_enabled() is True)
        with mock.patch.object(imagegen.LOGGER, "warning"):  # 失败日志不必刷屏
            check("接口不可达时静默返回 None", imagegen.generate_scene("market") is None)
        check("生图缓存落在 DATA_DIR/images", imagegen.cache_path("market", "p").parent == Path(config.DATA_DIR) / "images")
        check("未知主题提示词有兜底（不为空）", bool(imagegen.build_prompt("unknown-theme")))
    finally:
        for key, value in saved.items():
            if value is None:
                os.environ.pop(key, None)
            else:
                os.environ[key] = value

    # 5) 含插画的课件：图片形状 + AIGC 水印
    pkg = build_package({**base, "courseware": script})
    prs = Presentation(pkg["pptx"])
    pictures = [
        (index + 1, shape.name)
        for index, slide in enumerate(prs.slides)
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    ]
    check("场景卡插入图片形状（PICTURE）", bool(pictures), str(pictures[:3]))
    watermark = any(
        "AI 生成配图" in shape.text_frame.text
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame
    )
    check("AIGC 配图带「AI 生成配图」水印", watermark)

    # 6) 含图片的 deck 同样不得越界（裁切填充未破坏版面）
    EMU_IN = 914400
    width_in, height_in = prs.slide_width / EMU_IN, prs.slide_height / EMU_IN
    outside = [
        (index + 1, shape.shape_type, round(shape.left / EMU_IN, 2), round(shape.top / EMU_IN, 2))
        for index, slide in enumerate(prs.slides)
        for shape in slide.shapes
        if shape.left < 0 or shape.top < 0
        or shape.left + shape.width > prs.slide_width
        or shape.top + shape.height > prs.slide_height
    ]
    check("含插画课件形状不越出画布", not outside, str(outside[:3]))
    check("含插画课件仍为 16:9", abs(width_in / height_in - 16 / 9) < 0.01, f"{width_in:.2f}x{height_in:.2f}")

    # 7) 降级：素材库指向空目录仍能画出矢量场景（不抛异常、不残留图片）
    original_asset_dir = visuals.ASSET_DIR
    with tempfile.TemporaryDirectory() as empty_dir:
        visuals.ASSET_DIR = Path(empty_dir)
        try:
            degraded = build_package({**base, "courseware": script})
            prs2 = Presentation(degraded["pptx"])
            check("素材库缺失时降级为矢量场景（页数不变）", len(prs2.slides) == len(script), f"{len(prs2.slides)} vs {len(script)}")
            leftover = [
                shape.name
                for slide in prs2.slides
                for shape in slide.shapes
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
            ]
            check("降级后不含图片形状（走矢量）", not leftover, str(leftover[:3]))
            shapes_count = sum(len(slide.shapes) for slide in prs2.slides)
            check("降级后版面仍有矢量图形（不留白）", shapes_count > len(script) * 3, str(shapes_count))
        finally:
            visuals.ASSET_DIR = original_asset_dir

    # 8) _add_picture_fill：16:9 原图放进近方形框 → 居中裁切、不越界、可保存
    from PIL import Image as PILImage
    from pptx.util import Inches

    with tempfile.TemporaryDirectory() as tmp_dir:
        wide_path = Path(tmp_dir) / "wide.jpg"
        PILImage.new("RGB", (1600, 900), (228, 134, 43)).save(wide_path, quality=90)
        deck = Presentation()
        deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        shape = visuals._add_picture_fill(slide, wide_path, 1.0, 1.0, 3.0, 3.2)
        check("_add_picture_fill 返回图片形状", shape is not None)
        if shape is not None:
            check("16:9 原图进近方形框：左右对称裁切",
                  shape.crop_left > 0 and abs(shape.crop_left - shape.crop_right) < 1e-6,
                  f"{shape.crop_left:.3f}/{shape.crop_right:.3f}")
            check("裁切比例合法（左右之和 < 1）", shape.crop_left + shape.crop_right < 1.0,
                  f"{shape.crop_left + shape.crop_right:.3f}")
            check("图片严格占满目标框（不拉伸不越界）",
                  shape.left == Inches(1.0) and shape.width == Inches(3.0)
                  and shape.top + shape.height <= deck.slide_height,
                  f"left={shape.left} w={shape.width}")
            out_path = Path(tmp_dir) / "fill.pptx"
            deck.save(str(out_path))
            check("裁切填充后的 PPT 可正常保存", out_path.is_file() and out_path.stat().st_size > 0)

    # 9) 运行时素材目录不得残留大文件（src/ 仅存档，不参与解析）
    oversized_assets = [
        (item.name, item.stat().st_size)
        for item in visuals.ASSET_DIR.glob("*")
        if item.is_file() and item.stat().st_size > 400 * 1024
    ]
    check("运行时素材目录无 >400KB 文件（src/ 仅存档）", not oversized_assets, str(oversized_assets))
    check("AIGC 源图归档在 src/ 子目录", (visuals.ASSET_DIR / "src").is_dir())

    # 10) 水印裁切守卫：16:9 / 近方形 / 3:1 三种目标框都必须裁掉右、底
    #     承诺值唯一来源是 visuals.WATERMARK_GUARD（测试不写死魔法数），
    #     并用实测下限护栏把守卫值本身也锁住——否则把守卫调小（如回退到
    #     6%），"从常量派生"的断言会跟着变小、测试照样全绿，水印重新泄漏。
    guard = visuals.WATERMARK_GUARD
    source = visuals.ASSET_DIR / "market.jpg"
    if not source.is_file():
        found = visuals.scene_image("market")
        source = found if found is not None else source
    guard_failures = []
    if guard < WATERMARK_MIN_GUARD:
        guard_failures.append(
            f"守卫 {guard:.1%} 低于实测下限 {WATERMARK_MIN_GUARD:.1%}"
            f"（水印起点 {WATERMARK_MEASURED_START:.1%}W，6% 会切进「AI」字头）"
        )
    for label, box_w, box_h in (("16:9", 4.8, 2.7), ("近方形", 3.0, 3.2), ("3:1", 5.4, 1.8)):
        deck = Presentation()
        deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        pic = visuals._add_picture_fill(slide, source, 5.0, 2.0, box_w, box_h)
        if pic is None:
            guard_failures.append(f"{label}:未插入")
            continue
        inside = (
            pic.left >= 0 and pic.top >= 0
            and pic.left + pic.width <= deck.slide_width
            and pic.top + pic.height <= deck.slide_height
        )
        if pic.crop_right < guard - 1e-9:
            guard_failures.append(f"{label}：右裁 {pic.crop_right:.4f} < 承诺 {guard:.0%}")
        if pic.crop_bottom < guard - 1e-9:
            guard_failures.append(f"{label}：底裁 {pic.crop_bottom:.4f} < 承诺 {guard:.0%}")
        if not inside:
            guard_failures.append(f"{label}:越界")
        # 可见窗口比例须严格等于目标框比例（守卫不得引入变形）
        with PILImage.open(source) as raw:
            native_w, native_h = raw.size
        visible_ratio = ((1 - pic.crop_left - pic.crop_right) * native_w) / (
            (1 - pic.crop_top - pic.crop_bottom) * native_h
        )
        if abs(visible_ratio - box_w / box_h) > 1e-3:
            guard_failures.append(f"{label}:可见比例 {visible_ratio:.4f} ≠ {box_w / box_h:.4f}")
    check(f"水印守卫：右/底各裁 ≥{guard:.0%}、不越界、比例不变形", not guard_failures, str(guard_failures))

    # 11) png 兜底体积守卫：超大 png 自动降采样为 <400KB 的 JPEG 缓存
    source_png = visuals.ASSET_DIR / "src" / "market.png"
    if not source_png.is_file():
        archived = sorted((visuals.ASSET_DIR / "src").glob("*.png")) if (visuals.ASSET_DIR / "src").is_dir() else []
        source_png = archived[0] if archived else None
    if source_png is None:
        check("存在可复验的 png 源图（src/）", False, "src/*.png 缺失")
    else:
        original_asset_dir = visuals.ASSET_DIR
        with tempfile.TemporaryDirectory() as tmp_dir:
            big_dir = Path(tmp_dir) / "visuals"
            big_dir.mkdir()
            big_png = big_dir / "market.png"  # 主题键必须与查询一致
            big_png.write_bytes(source_png.read_bytes())
            cache_dir = Path(tmp_dir) / "data"
            big_size = big_png.stat().st_size
            visuals.ASSET_DIR = big_dir
            try:
                with mock.patch.object(config, "DATA_DIR", cache_dir), mock.patch.object(visuals.LOGGER, "warning"):
                    shrunk = visuals.scene_image("market")
            finally:
                visuals.ASSET_DIR = original_asset_dir
            check("超大 png 素材确实 >400KB（触发守卫）", big_size > 400 * 1024, f"{big_size // 1024}KB")
            check(
                "降采样后 <400KB 且落在 images 缓存目录",
                shrunk is not None and shrunk.is_file()
                and shrunk.stat().st_size < 400 * 1024
                and shrunk.parent == cache_dir / "images",
                f"{shrunk} {shrunk.stat().st_size // 1024 if shrunk and shrunk.is_file() else 0}KB",
            )


# ---------------------------------------------------------------------------
# storage
# ---------------------------------------------------------------------------
def test_storage() -> None:
    section("storage 存储层")
    from xhzh import storage

    params = {"grade": "七年级", "subject": "数学", "textbook": "人教版", "topic": "存储探针"}
    gid = storage.save_generation(params, status="running")
    check("插入返回 id", isinstance(gid, int) and gid > 0)

    storage.update_generation_status(gid, "done", detail={"rural_warnings": ["w1"]})
    storage.update_generation_status(gid, "done", detail={"kb_cards": ["K01"]})
    with sqlite3.connect(storage.config.DB_PATH) as conn:
        row = conn.execute("SELECT status, detail FROM generations WHERE id=?", (gid,)).fetchone()
    detail = json.loads(row[1])
    check("detail 合并而非覆盖", detail.get("rural_warnings") == ["w1"] and detail.get("kb_cards") == ["K01"])

    stale = storage.mark_stale_running_failed("测试清理")
    check("清理残留 running（≥0 且幂等可重复）", stale >= 0)
    stale2 = storage.mark_stale_running_failed("测试清理")
    check("重复清理无副作用", stale2 == 0, f"第二次清理了 {stale2} 条")

    rows = storage.list_recent_generations(5)
    check("历史查询返回元组", isinstance(rows, list) and len(rows) <= 5)


# ---------------------------------------------------------------------------
# app 前端
# ---------------------------------------------------------------------------
def test_app() -> None:
    section("app 前端层")
    import app

    out = next(app.generate("七年级", "数学", "人教版", "   "))
    check("空白课题被拒绝", "请输入课题" in out[2])

    events = list(app.generate("七年级", "语文", "人教版", "比喻拟人修辞"))
    final = events[-1]
    check("输出元组结构为 9", len(final) == 9)
    check("三件套路径有效", all(Path(p).exists() for p in final[4:7]))
    check("教案预览非空", "# 比喻拟人修辞" in final[1])
    check("生成明细含知识点卡", "知识点卡引用" in final[2])
    check("历史记录含本地时间", re.search(r"\d{2}-\d{2} \d{2}:\d{2}", final[3]) is not None, final[3][:80])

    # 交互防护：生成期间按钮禁用，结束后恢复
    first = events[0]
    check("生成中按钮被禁用", isinstance(first[7], dict) and first[7].get("interactive") is False, str(first[7]))
    check("结束后按钮恢复", isinstance(final[7], dict) and final[7].get("interactive") is True, str(final[7]))
    check("完成后下载区自动展开", isinstance(final[8], dict) and final[8].get("open") is True, str(final[8]))

    # 课题前后空格被规范化（不入文件名/数据库）
    out = list(app.generate("七年级", "数学", "人教版", "  一元一次方程  "))
    check("首尾空格课题可生成", out[-1][4] is not None)

    # 业务一致性：mock 教案与生成明细的课标口径不得互相矛盾
    consistent = list(app.generate("七年级", "数学", "人教版", "一元一次方程的实际应用"))[-1]
    check("有匹配时明细不含'无匹配'", "无匹配" not in consistent[2], consistent[2][:120])
    check("有匹配时明细含知识点卡 ID", "MAT-7A-3-0-K01" in consistent[2])
    check("教案与明细课标口径一致", ("暂无课标依据" in consistent[1]) == ("无匹配" in consistent[2]))

    # 失败路径：步骤条保留失败节点（不重置为全 pending）
    from xhzh import config as _cfg
    os.environ["XHZH_MODE"] = "adp"
    importlib.reload(_cfg)
    try:
        fe = list(app.generate("七年级", "数学", "人教版", "失败路径探针"))[-1]
        check("失败时步骤条标红失败节点", "step-failed" in fe[0])
        check("失败时按钮恢复可点", fe[7].get("interactive") is True)
        check("失败留痕含错误信息", "生成失败" in fe[2] and "adp" in fe[2])
    finally:
        os.environ.pop("XHZH_MODE")
        importlib.reload(_cfg)

    ui = app.build_ui()
    check("UI 可构建", ui is not None)


# ---------------------------------------------------------------------------
# main
# ---------------------------------------------------------------------------
GROUPS = {
    "config": test_config,
    "kb": test_kb,
    "llm": test_llm,
    "rural": test_rural,
    "workflow": test_workflow_mock,
    "workflow-llm": test_workflow_llm_pathological,
    "assembly": test_assembly,
    "visuals": test_visuals,
    "storage": test_storage,
    "app": test_app,
}

if __name__ == "__main__":
    # 测试环境隔离：剥离 .env 注入的运行配置（含 LLM 凭据），保证默认假设确定
    # （config 默认 mock、LLM 未配置）。用例内自行 set 的环境变量不受影响。
    from xhzh import config

    os.environ["XHZH_SKIP_DOTENV"] = "1"  # 防 reload 时 .env 变量回注
    for key in (
        "XHZH_MODE", "XHZH_MOCK", "LLM_BASE_URL", "LLM_API_KEY", "LLM_MODEL",
        "XHZH_IMAGE_BASE_URL", "XHZH_IMAGE_API_KEY", "XHZH_IMAGE_MODEL",
        "XHZH_IMAGE_ALLOW", "XHZH_IMAGE_TIMEOUT",
    ):
        os.environ.pop(key, None)
    importlib.reload(config)

    selected = sys.argv[1:] or list(GROUPS)
    print("星火智教 · 测试套件")
    for name in selected:
        GROUPS[name]()
    print(f"\n{'=' * 46}")
    print(f"总计 {PASS + FAIL} 项：✅ {PASS} 通过 / ❌ {FAIL} 失败")
    if FAILURES:
        print("\n失败明细：")
        for item in FAILURES:
            print(f"  - {item}")
    sys.exit(1 if FAIL else 0)
