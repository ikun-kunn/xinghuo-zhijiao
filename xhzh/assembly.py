"""教学包组装：教案 docx + 课件 pptx + 打包 zip。

Notes
-----
赛事版最小实现：结构化内容 → 单模板双件套。
配图节点在 mock 模式下为空列表，组装器自动跳过图片页。
"""

from __future__ import annotations

import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any

from . import config


def build_package(result: dict[str, Any]) -> dict[str, str]:
    """Assemble docx + pptx and zip them into one teaching package.

    Parameters
    ----------
    result : dict
        Final ``result`` of :func:`xhzh.workflow.run_generation`.

    Returns
    -------
    dict
        Paths: ``{"docx": ..., "pptx": ..., "zip": ...}``.
    """
    config.ensure_dirs()
    meta = result["meta"]
    # 微秒级时间戳 + 冲突自增，保证"产物不可变"（同秒/并发不覆盖）
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
    # 课题截断：Windows 单段文件名上限 255，超长会导致 Permission denied
    topic_part = str(meta.get("topic", ""))[:40]
    base = f"{meta['grade']}{meta['subject']}_{topic_part}_{stamp}"
    safe_base = "".join(c for c in base if c.isalnum() or c in "_-（）()").strip() or "teaching_package"

    docx_path, pptx_path, zip_path = _unique_paths(config.OUTPUT_DIR, safe_base)

    _build_docx(result, docx_path)
    _build_pptx(result, pptx_path)
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.write(docx_path, docx_path.name)
        zf.write(pptx_path, pptx_path.name)
    return {"docx": str(docx_path), "pptx": str(pptx_path), "zip": str(zip_path)}


def _unique_paths(output_dir: Path, stem: str) -> tuple[Path, Path, Path]:
    """Pick three output paths that do not collide with existing files.

    Parameters
    ----------
    output_dir : Path
        Target directory.
    stem : str
        Base file name without extension.

    Returns
    -------
    tuple of Path
        ``(docx, pptx, zip)`` paths guaranteed not to exist yet.
    """
    suffix = 0
    while True:
        name = stem if suffix == 0 else f"{stem}-{suffix}"
        paths = (
            output_dir / f"{name}.docx",
            output_dir / f"{name}.pptx",
            output_dir / f"{name}.zip",
        )
        if not any(path.exists() for path in paths):
            return paths
        suffix += 1


def _build_docx(result: dict[str, Any], path: Path) -> None:
    """Write the lesson-plan Word file.

    Parameters
    ----------
    result : dict
        Workflow result.
    path : Path
        Output docx path.

    Returns
    -------
    None
    """
    from docx import Document
    from .workflow import render_plan_text

    document = Document()
    for line in render_plan_text(result).splitlines():
        if line.startswith("# "):
            document.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=2)
        elif line.startswith("### "):
            document.add_heading(line[4:], level=3)
        elif line.startswith("- "):
            document.add_paragraph(line[2:], style="List Bullet")
        elif line.startswith("> "):
            document.add_paragraph(line[2:])
        elif line.strip():
            document.add_paragraph(line)
    document.save(str(path))


def _build_pptx(result: dict[str, Any], path: Path) -> None:
    """Write the courseware PPT file (student-facing when a script exists).

    Parameters
    ----------
    result : dict
        Workflow result. When ``result["courseware"]`` holds a slide script
        (see :mod:`xhzh.courseware`), the engaging card-style renderer is
        used; otherwise the structured renderer runs as a fallback.
    path : Path
        Output pptx path.

    Returns
    -------
    None
    """
    script = [s for s in (result.get("courseware") or []) if isinstance(s, dict)]
    if script:
        _render_engaging(result, script, path)
    else:
        _render_structured(result, path)


def _render_structured(result: dict[str, Any], path: Path) -> None:
    """Write the structured (teacher-facing) PPT as a fallback.

    Parameters
    ----------
    result : dict
        Workflow result.
    path : Path
        Output pptx path.

    Returns
    -------
    None

    Notes
    -----
    页面结构对齐教案实施（评估报告 B 方案）：封面 → 学习目标 → 重难点 →
    教学流程 → 逐环节活动页 → 分层练习（每页≤2 题）→ 参考答案 → 板书 →
    小结与作业。练习与答案分离，避免单页文字墙。
    """
    from pptx import Presentation
    from pptx.util import Pt

    from .courseware import as_str_list as _as_str_list

    meta = result["meta"]
    stages = [
        stage for stage in (result.get("stages") or [])
        if isinstance(stage, dict) and str(stage.get("name", "")).strip()
    ]
    basic = _as_str_list(result.get("exercises_basic"))
    extend = _as_str_list(result.get("exercises_extend"))
    answers = _as_str_list(result.get("answers"))

    slides: list[tuple[str, list[str]]] = []

    # 1. 封面（副标题：年级/学科/教材 + 品牌）
    slides.append((
        str(meta["topic"]),
        [f"{meta['grade']} {meta['subject']} · {meta['textbook']}", "星火智教 · AI 生成"],
    ))
    # 2. 学习目标（分层，逐条短句展示）
    goal_lines = [f"基础 · {g}" for g in _as_str_list(result.get("goals_basic"))]
    goal_lines += [f"拓展 · {g}" for g in _as_str_list(result.get("goals_extend"))]
    slides.append(("学习目标", goal_lines or ["（待补充）"]))
    # 3. 重难点
    heavy_lines = [f"重点：{t}" for t in _as_str_list(result.get("key_points"))]
    heavy_lines += [f"难点：{t}" for t in _as_str_list(result.get("difficult_points"))]
    slides.append(("重难点", heavy_lines or ["（待补充）"]))
    # 4. 教学流程（环节 + 时长 + 总课时）
    if stages:
        total = sum(int(stage.get("minutes") or 0) for stage in stages)
        flow = [
            f"{index + 1}. {stage['name']}"
            + (f"（{stage['minutes']} 分钟）" if stage.get("minutes") else "")
            for index, stage in enumerate(stages)
        ]
        slides.append((f"教学流程（共 {total} 分钟）", flow))
        # 5.. 逐环节活动页（教学过程主体，支撑 45 分钟实施）
        for stage in stages:
            title = stage["name"] + (f"（{stage['minutes']} 分钟）" if stage.get("minutes") else "")
            desc = str(stage.get("desc", "")).strip()
            slides.append((str(title), [desc] if desc else ["教师引导学生完成本环节活动"]))
    # 练习分页：每页最多 2 题（避免文字墙）
    for index in range(0, len(basic), 2):
        slides.append((f"基础练习 {index // 2 + 1}", basic[index:index + 2]))
    for index in range(0, len(extend), 2):
        slides.append((f"拓展练习 {index // 2 + 1}", extend[index:index + 2]))
    # 参考答案（与题目分离展示）
    for index in range(0, len(answers), 3):
        slides.append((f"参考答案 {index // 3 + 1}", answers[index:index + 3]))
    # 板书设计
    slides.append(("板书设计", [result.get("board") or "（见教案）"]))
    # 小结与作业
    slides.append((
        "课堂小结与作业",
        [
            "小结：回顾建模步骤 —— 设未知数 → 找等量关系 → 列方程 → 检验作答",
            "作业：基础题 2 道必做，拓展题 1 道选做",
        ],
    ))

    presentation = Presentation()
    for index, (title, bullets) in enumerate(slides):
        layout = presentation.slide_layouts[0 if index == 0 else 1]
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text = title
        if index == 0:
            # 首页副标题（年级/学科/教材/品牌）：此前直接 continue 导致副标题为空
            if len(slide.placeholders) > 1:
                slide.placeholders[1].text = "\n".join(bullets)
            continue
        body = slide.placeholders[1].text_frame
        body.clear()
        for i, bullet in enumerate(bullets):
            paragraph = body.paragraphs[0] if i == 0 else body.add_paragraph()
            paragraph.text = bullet
            # 题干较长，22pt 兼顾可读性与版面；纯短句页同样适用
            paragraph.font.size = Pt(22 if len(bullet) > 24 else 24)
    presentation.save(str(path))


# ----------------------------------------------------------------------
# 学生向课件渲染：卡片式版式
# 背景色 → 标题色带 + 环节标签 → 引导句 → 内容卡 → 教师提示条 → 页码
# 全部基于 Blank 版式手绘，不依赖模板占位符，县乡学校旧版 PowerPoint 亦可读。
# ----------------------------------------------------------------------

_W_IN = 13.333  # 16:9 画布宽（英寸）
_H_IN = 7.5
_M = 0.6  # 页边距

_BG = "FDF7EF"
_INK = "2C2A26"
_GREEN = "2E7D5B"
_ORANGE = "E4862B"
_BLUE = "2F6FB5"
_CARD = "FFFFFF"
_CARD_LINE = "EFE3D2"
_TIP_BG = "FFF3E3"
_TIP_INK = "8A6A3F"
_MUTED = "9A9088"
# 中文渲染字体：县乡学校机器常见，且比宋体更适合投影
_FONT = "微软雅黑"

# 页型 → (主色, 浅色底)
_KIND_STYLE = {
    "cover": (_GREEN, "E7F1EB"),
    "hook": (_BLUE, "E8F0F9"),
    "quest": (_GREEN, "E7F1EB"),
    "explore": (_GREEN, "E7F1EB"),
    "route": (_ORANGE, "FCEEDD"),
    "practice": (_ORANGE, "FCEEDD"),
    "game": (_ORANGE, "FCEEDD"),
    "answer": (_GREEN, "E7F1EB"),
    "board": (_BLUE, "E8F0F9"),
    "summary": (_GREEN, "E7F1EB"),
}
# 页型 → 默认环节标签（脚本未给 tag 时兜底）
_KIND_TAG = {
    "hook": "第 1 关",
    "quest": "任务清单",
    "explore": "锦囊",
    "route": "路线图",
    "practice": "基础关",
    "game": "终极关",
    "answer": "揭晓",
    "board": "方法卡",
    "summary": "通关",
}
# 需要给卡片编号标签的页型
_CARD_LABEL = {"practice": "关卡", "game": "终极关", "answer": "答案"}


def _render_engaging(result: dict[str, Any], script: list[dict[str, Any]], path: Path) -> None:
    """Render the student-facing courseware PPT from a slide script.

    Parameters
    ----------
    result : dict
        Workflow result (used for cover metadata).
    script : list of dict
        Slide specs from :func:`xhzh.courseware.build_script`.
    path : Path
        Output pptx path.

    Returns
    -------
    None

    Notes
    -----
    每页先经 :func:`xhzh.visuals.plan_visual` 规划视觉元素（场景插图 /
    图标徽章 / 数量关系图），再按页型融入版面：封面与练习页占右侧栏，
    任务/锦囊/方法卡在卡片内嵌徽章，answer/route/summary 保持纯文字。
    """
    from pptx import Presentation
    from pptx.util import Inches

    from . import visuals

    presentation = Presentation()
    presentation.slide_width = Inches(_W_IN)
    presentation.slide_height = Inches(_H_IN)
    layout = _blank_layout(presentation)
    total = len(script)
    for index, spec in enumerate(script):
        if spec.get("kind") == "cover":
            _draw_cover(presentation, layout, result, spec, visuals)
        else:
            _draw_content(presentation, layout, spec, index + 1, total, visuals)
    presentation.save(str(path))


def _draw_cover(presentation, layout, result: dict[str, Any], spec: dict[str, Any], visuals=None) -> None:
    """Draw the half-panel cover slide with an optional theme scene.

    Parameters
    ----------
    presentation : pptx.Presentation
    layout : pptx.slide.SlideLayout
        Blank layout.
    result : dict
        Workflow result.
    spec : dict
        Cover slide spec.
    visuals : module or None
        :mod:`xhzh.visuals`; ``None`` skips the scene (defensive default).

    Returns
    -------
    None
    """
    slide = presentation.slides.add_slide(layout)
    _rect(slide, 0, 0, _W_IN, _H_IN, _BG)
    _rect(slide, 0, 0, 5.15, _H_IN, _GREEN)
    meta = result.get("meta", {})
    title = str(spec.get("title") or meta.get("topic") or "本课课件")
    size = 34 if len(title) <= 12 else 28
    _text(slide, 0.78, 1.9, 3.9, 2.2, title, size=size, bold=True, color="FFFFFF")
    subtitle = str((spec.get("points") or [""])[0]).strip()
    if subtitle:
        _text(slide, 0.78, 4.2, 3.9, 0.9, subtitle, size=14, color="D7EBE1")
    _text(slide, 0.78, 5.15, 3.9, 0.5, "星火智教 · AI 生成课件", size=12, color="B9D8CA")

    lead = str(spec.get("lead") or "").strip()
    if lead:
        _rect(slide, 5.85, 1.8, 0.16, 1.5, _ORANGE)
        _text(slide, 6.3, 1.8, 6.4, 1.5, lead, size=22, bold=True, color=_INK, anchor="middle")
    tag = str(spec.get("tag") or "课堂闯关").strip()
    _round(slide, 6.3, 3.55, 2.3, 0.62, "FCEEDD")
    _text(slide, 6.3, 3.55, 2.3, 0.62, tag, size=14, bold=True, color=_ORANGE, align="center", anchor="middle")
    for offset, color in enumerate((_ORANGE, _BLUE, _GREEN)):
        _rect(slide, 6.3 + offset * 0.58, 4.5, 0.42, 0.16, color)
    # 主题场景插图（右下区域）：模型规划或按课题关键词推断
    if visuals is not None:
        points = [str(p) for p in (spec.get("points") or []) if str(p).strip()]
        planned = visuals.plan_visual(spec, "cover", points or [title])
        if planned and planned.get("type") == "scene":
            visuals.draw_cover_scene(slide, planned.get("theme") or "market", 6.3, 4.95, 5.6, 2.0)
    _rect(slide, 0, _H_IN - 0.26, _W_IN, 0.26, _ORANGE)


def _draw_content(presentation, layout, spec: dict[str, Any], number: int, total: int, visuals=None) -> None:
    """Draw one content slide (everything except the cover).

    Parameters
    ----------
    presentation : pptx.Presentation
    layout : pptx.slide.SlideLayout
        Blank layout.
    spec : dict
        Slide spec with ``kind`` / ``title`` / ``lead`` / ``points`` / ``tip`` / ``tag``.
    number : int
        1-based page number.
    total : int
        Total page count.
    visuals : module or None
        :mod:`xhzh.visuals`; ``None`` renders text-only pages.

    Returns
    -------
    None
    """
    slide = presentation.slides.add_slide(layout)
    kind = str(spec.get("kind") or "explore")
    color, soft = _KIND_STYLE.get(kind, (_GREEN, "E7F1EB"))

    _rect(slide, 0, 0, _W_IN, _H_IN, _BG)
    _round(slide, _M, 0.42, 7.9, 0.92, color)
    _text(
        slide, _M + 0.3, 0.42, 7.4, 0.92, str(spec.get("title") or ""),
        size=26, bold=True, color="FFFFFF", anchor="middle",
    )
    tag = str(spec.get("tag") or _KIND_TAG.get(kind, "")).strip()
    if tag:
        _round(slide, _W_IN - _M - 3.0, 0.5, 3.0, 0.76, soft)
        _text(
            slide, _W_IN - _M - 3.0, 0.5, 3.0, 0.76, tag,
            size=14, bold=True, color=color, align="center", anchor="middle",
        )
    _rect(slide, _M, 1.5, _W_IN - 2 * _M, 0.045, _CARD_LINE)

    lead = str(spec.get("lead") or "").strip()
    if lead:
        _text(slide, _M + 0.08, 1.68, _W_IN - 2 * _M - 0.16, 0.6, lead, size=19, bold=True, color=color)

    points = [p for p in (spec.get("points") or []) if str(p).strip()]

    # 视觉元素规划：练习/悬念页右侧栏放图表或场景，任务/锦囊/方法卡嵌图标徽章
    side_visual = None
    icons: list[str] | None = None
    if visuals is not None:
        planned = visuals.plan_visual(spec, kind, points)
        if planned:
            if planned["type"] in ("scene", "diagram") and kind in ("hook", "practice", "game"):
                side_visual = planned
            elif planned["type"] == "iconrow":
                icons = planned["icons"]

    if kind == "route":
        _draw_route(slide, points, color)
    elif side_visual is not None:
        # 左侧内容卡收窄，右侧留给视觉元素
        card_width = 8.05
        _draw_cards(slide, points, color, kind, width=card_width, icons=None)
        visuals.draw_side_visual(
            slide, side_visual,
            _M + card_width + 0.25, 2.42,
            _W_IN - _M - (_M + card_width + 0.25), _H_IN - 1.16 - 2.42,
            color,
        )
    else:
        _draw_cards(slide, points, color, kind, icons=icons)

    tip = str(spec.get("tip") or "").strip()
    if tip:
        _round(slide, _M, _H_IN - 0.98, _W_IN - 2 * _M, 0.68, _TIP_BG)
        _text(
            slide, _M + 0.25, _H_IN - 0.98, _W_IN - 2 * _M - 0.5, 0.68,
            f"教师提示 · {tip}", size=12, color=_TIP_INK, anchor="middle",
        )
    _text(
        slide, _W_IN - _M - 1.4, _H_IN - 0.44, 1.4, 0.3,
        f"{number} / {total}", size=10, color=_MUTED, align="right",
    )


def _draw_cards(slide, points: list[str], color: str, kind: str, width: float | None = None, icons: list[str] | None = None) -> None:
    """Draw vertical stacked content cards, optionally with icon badges.

    Parameters
    ----------
    slide : pptx.slide.Slide
    points : list of str
        Card texts; empty input renders one placeholder card.
    color : str
        Accent hex color.
    kind : str
        Slide kind; ``practice`` / ``game`` / ``answer`` get numbered chips.
    width : float or None
        Card width in inches; defaults to the full content width (narrowed
        when a side visual occupies the right region).
    icons : list of str or None
        Icon names (one per card) drawn as badges on the card's right.

    Returns
    -------
    None
    """
    if not points:
        points = ["（内容待补充）"]
    count = len(points)
    top, bottom, gap = 2.42, _H_IN - 1.16, 0.2
    available = bottom - top
    height = min(1.5, max(0.7, (available - gap * (count - 1)) / count))
    start = top + (available - (height * count + gap * (count - 1))) / 2
    if width is None:
        width = _W_IN - 2 * _M
    label = _CARD_LABEL.get(kind)
    for index, point in enumerate(points):
        top_y = start + index * (height + gap)
        _round(slide, _M, top_y, width, height, _CARD, line=_CARD_LINE)
        _rect(slide, _M + 0.02, top_y + 0.08, 0.09, height - 0.16, color)
        text_top, text_height = top_y + 0.06, height - 0.12
        text_width = width - 0.68
        if label:
            _round(slide, _M + 0.3, top_y + 0.12, 1.5, 0.36, color)
            _text(
                slide, _M + 0.3, top_y + 0.12, 1.5, 0.36, f"{label} {index + 1}",
                size=11, bold=True, color="FFFFFF", align="center", anchor="middle",
            )
            text_top, text_height = top_y + 0.54, height - 0.62
        # 图标徽章：卡片右侧（预留徽章位后正文收窄）
        badge = icons[index] if icons and index < len(icons) else None
        from . import visuals as _visuals
        if badge:
            badge_size = 0.44
            badge_x = _M + width - badge_size - 0.18
            badge_y = top_y + (height - badge_size) / 2
            _visuals.draw_icon_badge(slide, badge, badge_x, badge_y, color, size=badge_size)
            text_width = width - 0.36 - (badge_size + 0.3)
        text = str(point)
        size = 19 if len(text) <= 26 else (17 if len(text) <= 40 else 15)
        _text(
            slide, _M + 0.36, text_top, text_width, text_height, text,
            size=size, color=_INK, anchor="middle",
        )


def _draw_route(slide, points: list[str], color: str) -> None:
    """Draw a horizontal step-flow (numbered chips with arrows).

    Parameters
    ----------
    slide : pptx.slide.Slide
    points : list of str
        Step texts; at most four are drawn.
    color : str
        Accent hex color.

    Returns
    -------
    None
    """
    steps = [p for p in points][:4]
    if not steps:
        return
    gap = 0.34
    width = (_W_IN - 2 * _M - gap * (len(steps) - 1)) / len(steps)
    height, top_y = 2.1, 3.0
    for index, step in enumerate(steps):
        left = _M + index * (width + gap)
        _round(slide, left, top_y, width, height, _CARD, line=_CARD_LINE)
        diameter = 0.62
        _oval(slide, left + width / 2 - diameter / 2, top_y + 0.26, diameter, diameter, color)
        _text(
            slide, left + width / 2 - diameter / 2, top_y + 0.26, diameter, diameter, str(index + 1),
            size=20, bold=True, color="FFFFFF", align="center", anchor="middle",
        )
        text = str(step)
        size = 17 if len(text) <= 18 else 15
        _text(
            slide, left + 0.22, top_y + 1.02, width - 0.44, height - 1.15, text,
            size=size, color=_INK, align="center", anchor="middle",
        )
        if index < len(steps) - 1:
            _arrow(slide, left + width + 0.05, top_y + height / 2 - 0.13, 0.26, 0.26, color)


# ----------------------------------------------------------------------
# 绘制原语
# ----------------------------------------------------------------------
def _blank_layout(presentation):
    """Return the blank layout of a presentation.

    Parameters
    ----------
    presentation : pptx.Presentation

    Returns
    -------
    pptx.slide.SlideLayout
        Layout named ``"Blank"``, or index 6 as a fallback.
    """
    for layout in presentation.slide_layouts:
        if layout.name == "Blank":
            return layout
    return presentation.slide_layouts[6]


def _paint(shape, fill: str, line: str | None = None) -> None:
    """Apply solid fill and optional outline to an autoshape.

    Parameters
    ----------
    shape : pptx.shapes.autoshape.Shape
    fill : str
        Hex RGB fill color.
    line : str or None
        Hex RGB outline color; ``None`` removes the outline.

    Returns
    -------
    None
    """
    from pptx.dml.color import RGBColor
    from pptx.util import Pt

    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        shape.line.color.rgb = RGBColor.from_string(line)
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False


def _rect(slide, x: float, y: float, w: float, h: float, fill: str):
    """Add a rectangle shape.

    Parameters
    ----------
    slide : pptx.slide.Slide
    x, y, w, h : float
        Geometry in inches.
    fill : str
        Hex RGB fill color.

    Returns
    -------
    pptx.shapes.autoshape.Shape
    """
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    _paint(shape, fill)
    return shape


def _round(slide, x: float, y: float, w: float, h: float, fill: str, line: str | None = None):
    """Add a rounded rectangle (cards, title bar, chips).

    Parameters
    ----------
    slide : pptx.slide.Slide
    x, y, w, h : float
        Geometry in inches.
    fill : str
        Hex RGB fill color.
    line : str or None
        Hex RGB outline color.

    Returns
    -------
    pptx.shapes.autoshape.Shape
    """
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    # 圆角比例：小尺寸（标签）用大圆角，大卡片用小圆角，避免"药丸"感
    shape.adjustments[0] = 0.35 if min(w, h) < 0.5 else 0.06
    _paint(shape, fill, line)
    return shape


def _oval(slide, x: float, y: float, w: float, h: float, fill: str):
    """Add an oval (step number badge).

    Parameters
    ----------
    slide : pptx.slide.Slide
    x, y, w, h : float
        Geometry in inches.
    fill : str
        Hex RGB fill color.

    Returns
    -------
    pptx.shapes.autoshape.Shape
    """
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
    _paint(shape, fill)
    return shape


def _arrow(slide, x: float, y: float, w: float, h: float, fill: str):
    """Add a right arrow connector between steps.

    Parameters
    ----------
    slide : pptx.slide.Slide
    x, y, w, h : float
        Geometry in inches.
    fill : str
        Hex RGB fill color.

    Returns
    -------
    pptx.shapes.autoshape.Shape
    """
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    shape = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    _paint(shape, fill)
    return shape


def _text(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    size: int = 18,
    bold: bool = False,
    color: str = _INK,
    align: str = "left",
    anchor: str = "top",
):
    """Add a text box with Chinese-safe font settings.

    Parameters
    ----------
    slide : pptx.slide.Slide
    x, y, w, h : float
        Geometry in inches.
    text : str
        Content; ``\\n`` creates additional paragraphs.
    size : int
        Font size in points.
    bold : bool
    color : str
        Hex RGB text color.
    align : str
        ``"left"`` / ``"center"`` / ``"right"``.
    anchor : str
        ``"top"`` / ``"middle"`` / ``"bottom"``.

    Returns
    -------
    pptx.shapes.autoshape.Shape
        The text box shape.
    """
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches

    aligns = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
    anchors = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchors.get(anchor, MSO_ANCHOR.TOP)
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    for index, line in enumerate(str(text).split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = aligns.get(align, PP_ALIGN.LEFT)
        run = paragraph.add_run()
        run.text = line
        _set_run_font(run, size, bold, color)
    return box


def _set_run_font(run, size: int, bold: bool, color: str) -> None:
    """Set size / bold / color / East-Asian typeface on one run.

    Parameters
    ----------
    run : pptx.text.text._Run
    size : int
        Font size in points.
    bold : bool
    color : str
        Hex RGB text color.

    Returns
    -------
    None

    Notes
    -----
    ``font.name`` 只写 latin 字体，中文在 PowerPoint 上会回退到默认宋体；
    此处额外写入 ``a:ea``，保证中文按「微软雅黑」渲染。
    """
    from pptx.dml.color import RGBColor
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import nsdecls, qn
    from pptx.util import Pt

    font = run.font
    font.size = Pt(size)
    font.bold = bold
    font.color.rgb = RGBColor.from_string(color)
    font.name = _FONT
    rPr = run._r.get_or_add_rPr()
    east_asian = rPr.find(qn("a:ea"))
    if east_asian is None:
        east_asian = parse_xml(f'<a:ea {nsdecls("a")}/>')
        rPr.append(east_asian)
    east_asian.set("typeface", _FONT)
